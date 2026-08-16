from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SKILL_ROOT = Path(
    "/Users/onnwee/.codex/plugins/cache/openai-primary-runtime/documents/26.709.11516/skills/documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


OUTPUT = Path("/Users/onnwee/Projects/subcult/hasanara/HasAnAra-Frontend-Quality-Review.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "172B4D"
MUTED = "5F6B7A"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
RED = "9B1C1C"
RED_FILL = "FDECEC"
GOLD = "7A5A00"
GREEN = "1F6B45"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, *, color: str = "D9DEE7", size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_run_font(run, name: str = "Calibri", size: float | None = None, *, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def set_paragraph_keep(paragraph, *, keep_next=False, keep_lines=True, widow=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for tag, enabled in (("keepNext", keep_next), ("keepLines", keep_lines), ("widowControl", widow)):
        existing = p_pr.find(qn(f"w:{tag}"))
        if enabled and existing is None:
            p_pr.append(OxmlElement(f"w:{tag}"))


def add_callout(doc: Document, label: str, text: str, *, color=RED, fill=RED_FILL):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.line_spacing = 1.1
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    borders.append(left)
    p_pr.append(borders)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=11, bold=True, color=color)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color=INK)
    return paragraph


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None):
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        run = p.add_run(text)
        set_run_font(run, color=INK)
    set_paragraph_keep(p)
    return p


def add_bullet(doc: Document, text: str, *, level: int = 0, bold_prefix: str | None = None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_run_font(first, bold=True, color=INK)
        rest = p.add_run(text[len(bold_prefix):])
        set_run_font(rest, color=INK)
    else:
        set_run_font(p.add_run(text), color=INK)
    set_paragraph_keep(p)
    return p


def add_numbered(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    set_run_font(p.add_run(text), color=INK)
    set_paragraph_keep(p)
    return p


def add_code_ref(doc: Document, path: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(path)
    set_run_font(run, name="Courier New", size=9.2, color=DARK_BLUE)
    set_paragraph_keep(p)
    return p


def add_issue(doc: Document, title: str, severity: str, summary: str, refs: list[str], fixes: list[str]):
    heading = doc.add_paragraph(style="Heading 2")
    tag = heading.add_run(f"{severity}  ")
    set_run_font(tag, size=12.5, bold=True, color=RED if severity in {"CRITICAL", "P1"} else GOLD)
    title_run = heading.add_run(title)
    set_run_font(title_run, size=13, bold=True, color=BLUE)
    set_paragraph_keep(heading, keep_next=True)
    add_body(doc, summary)
    if refs:
        label = doc.add_paragraph()
        label.paragraph_format.space_after = Pt(2)
        set_run_font(label.add_run("Affected code"), size=10, bold=True, color=MUTED)
        set_paragraph_keep(label, keep_next=True)
        for ref in refs:
            add_code_ref(doc, ref)
    if fixes:
        label = doc.add_paragraph()
        label.paragraph_format.space_before = Pt(4)
        label.paragraph_format.space_after = Pt(2)
        set_run_font(label.add_run("Recommended action"), size=10, bold=True, color=MUTED)
        set_paragraph_keep(label, keep_next=True)
        for fix in fixes:
            add_bullet(doc, fix)


def add_heading(doc: Document, text: str, level: int = 1, *, page_break=False):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    if page_break:
        p.paragraph_format.page_break_before = True
    set_paragraph_keep(p, keep_next=True)
    return p


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    set_run_font(hp.add_run("HASANARA  /  FRONTEND QUALITY REVIEW"), size=8.5, bold=True, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    set_run_font(fp.add_run("Page "), size=9, color=MUTED)
    add_page_field(fp, "PAGE")
    set_run_font(fp.add_run(" of "), size=9, color=MUTED)
    add_page_field(fp, "NUMPAGES")


def add_metadata_table(doc: Document):
    rows = [
        ("Verdict", "Not production-ready without fixes"),
        ("Scope", "Current checkout at 873a78b; React frontend, relevant API contracts, tests, CI, and documentation"),
        ("Strongest signal", "Build and unit tests pass; security, quality-gate, browser-coverage, and contract issues remain"),
        ("Review mode", "Read-only repository review; no source files or services changed"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Review field"
    table.rows[0].cells[1].text = "Assessment"
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10, bold=True, color=DARK_BLUE)
    for row in table.rows[1:]:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=10, bold=index == 0, color=INK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [1900, 7460], table_width_dxa=9360, indent_dxa=120)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_verification_table(doc: Document):
    data = [
        ("Unit tests", "PASS", "23 files; 95 passed, 2 skipped"),
        ("Coverage thresholds", "PASS", "69.4% lines; 67.31% statements; 57.19% branches"),
        ("TypeScript", "PASS", "npx tsc --noEmit"),
        ("Production build", "PASS", "460.42 KB JS / 94.14 KB CSS"),
        ("ESLint", "FAIL", "20 errors and 1 warning"),
        ("Prettier", "FAIL", "42 files require formatting"),
        ("git diff --check", "PASS", "No whitespace errors"),
        ("Browser E2E", "NOT RUN", "Docker-starting suite skipped; static inspection shows route and contract drift"),
        ("Dependency audit", "ACTION", "Production and development dependency advisories require triage"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ("Check", "Result", "Evidence")):
        cell.text = value
    for check, result, evidence in data:
        cells = table.add_row().cells
        cells[0].text = check
        cells[1].text = result
        cells[2].text = evidence
    for cell in table.rows[0].cells:
        set_cell_shading(cell, LIGHT_GRAY)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=10, bold=True, color=DARK_BLUE)
    for row in table.rows[1:]:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx == 1:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cell.paragraphs[0].runs:
                result = row.cells[1].text
                color = GREEN if result == "PASS" else RED if result == "FAIL" else GOLD
                set_run_font(run, size=9.5, bold=idx == 1, color=color if idx == 1 else INK)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [2450, 1450, 5460], table_width_dxa=9360, indent_dxa=120)


def build_document() -> None:
    doc = Document()
    configure_styles(doc)
    configure_page(doc)

    # Title block: memo_masthead pattern resolved through the standard_business_brief preset.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(18)
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("TECHNICAL QUALITY REVIEW"), size=9.5, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("HasAnAra Web Frontend"), size=25, bold=True, color=INK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(13)
    set_run_font(
        subtitle.add_run("Code quality, functionality, performance, accessibility, scalability, and documentation assessment"),
        size=13,
        color=MUTED,
    )

    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(15)
    set_run_font(meta.add_run(f"Prepared {date.today().strftime('%B %-d, %Y')}  |  Commit 873a78b  |  Branch main"), size=9.5, bold=True, color=MUTED)

    add_callout(
        doc,
        "Verdict",
        "Not production-ready without fixes. The archive experience is promising and the production build succeeds, but a serious injection risk, broken or incomplete flows, stale browser coverage, failing quality gates, and documentation drift must be addressed.",
    )
    add_metadata_table(doc)

    add_heading(doc, "Executive Summary", 1)
    add_body(
        doc,
        "The frontend has a strong product foundation: archive-wide search, topic intelligence, timestamp deep links, transcript playback synchronization, citation-backed chapters, saved moments, and responsive discovery views. The most important risks are concentrated in trust boundaries and contract integrity rather than in basic compilation.",
    )
    add_bullet(doc, "Critical: external transcript/search content is inserted as raw HTML in four user-facing surfaces.")
    add_bullet(doc, "Important: anonymous saving, saved-search filters, authentication failures, admin access feedback, route errors, and VOD-to-VOD player transitions have broken or misleading behavior.")
    add_bullet(doc, "Delivery risk: unit tests and build pass, while ESLint, Prettier, browser specifications, and dependency audits show unresolved issues.")
    add_bullet(doc, "Scalability risk: all routes are eagerly bundled and long transcript views rerender continuously during playback.")
    add_bullet(doc, "Documentation risk: product, architecture, accessibility, PWA, design-system, and test documentation no longer describe the shipped frontend accurately.")

    add_heading(doc, "Review Triage", 1)
    add_bullet(doc, "Docs-only review: No")
    add_bullet(doc, "React performance review: Yes")
    add_bullet(doc, "UI guidelines audit: Yes")
    add_body(doc, "The project is a React 19 single-page application containing large transcript views, route-level data fetching, extensive interactive controls, and custom accessibility behavior. UI findings were assessed against the current Vercel Web Interface Guidelines.")

    add_heading(doc, "Strengths", 1)
    for item in [
        "The production build and TypeScript type-check succeed.",
        "All 23 unit-test files pass: 95 tests passed and 2 are skipped.",
        "Configured coverage thresholds pass, with 69.4% line coverage.",
        "Search and VOD-library state is generally represented in the URL.",
        "The interface has a coherent editorial identity, semantic color tokens, a skip link, global focus treatment, reduced-motion CSS, and responsive layouts.",
        "Transcript chapters, active playback tracking, timestamp links, saved moments, topic statistics, Explore periods, and archive metadata form a compelling product foundation.",
        "Backend admin endpoints are authorization-gated even though the frontend admin guard needs improvement.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Critical Security Finding", 1)
    add_issue(
        doc,
        "Search and transcript snippets permit stored HTML injection",
        "CRITICAL",
        "Search snippets from PostgreSQL/OpenSearch and transcript data are rendered through dangerouslySetInnerHTML without sanitization. The backend passes external transcript text through ts_headline or OpenSearch highlighting but does not escape arbitrary HTML before returning it. A malicious tag or event attribute could therefore become stored cross-site scripting.",
        [
            "frontend/src/components/archive/SearchMomentsList.tsx:50",
            "frontend/src/components/archive/TopicMentionCard.tsx:21",
            "frontend/src/routes/TopicPage.tsx:191",
            "frontend/src/components/video/PlainTranscriptTurns.tsx:89",
        ],
        [
            "Preferred: return plain text plus structured highlight ranges and render React nodes.",
            "Near-term: sanitize with a strict allowlist permitting only <mark>, and configure every search backend to emit the same tag.",
            "Add hostile-snippet regression tests covering event attributes, script tags, malformed markup, and safe highlighting.",
        ],
    )
    add_callout(
        doc,
        "Contract mismatch",
        "PostgreSQL normally emits <b>, OpenSearch defaults to <em>, API documentation shows <em>, while frontend styling and tests expect <mark>. Fixing the security boundary should also establish one explicit highlighting contract.",
        color=GOLD,
        fill="FFF8E1",
    )

    add_heading(doc, "Important Functional Findings", 1)
    add_issue(
        doc,
        "Anonymous users can save moments but cannot open their saved collection",
        "P1",
        "Search and Video pages write anonymous saves to localStorage, and FavoritesPage contains a local-mode UI. Both /saved and /favorites are wrapped in Protected, so anonymous users see a sign-in requirement rather than their saved moments. This contradicts the page copy that local moments still work in the browser.",
        ["frontend/src/main.tsx:39-53", "frontend/src/routes/FavoritesPage.tsx:102-138"],
        [
            "Make the Saved page accessible anonymously and gate only synchronized server data.",
            "Offer migration of local moments after login.",
            "Add an anonymous save-to-reopen browser test.",
        ],
    )
    add_issue(
        doc,
        "Supported search filters are silently discarded",
        "P1",
        "The URL parser and saved-search model support source, category, duration, sort, video, limit, and offset. SearchPage then replaces source, category, duration, and sort with undefined. Reopened saved searches and deep links can therefore return broader or differently ordered results than requested.",
        ["frontend/src/features/search/filters.ts:5-34", "frontend/src/routes/SearchPage.tsx:150-160"],
        [
            "Wire every supported filter into the request and user interface, or remove unsupported fields from frontend contracts until implemented.",
            "Add an integration test that saves and reopens a multi-filter query.",
        ],
    )
    add_issue(
        doc,
        "Async route requests can overwrite newer state",
        "P1",
        "Search, Topic, Streams, Video, and Explore issue requests without cancellation or request-generation guards. A slow response for query, period, or video A can arrive after B and replace the current display. Search also retries the flat endpoint after any grouped-search failure, masking authorization, quota, timeout, or server errors and doubling load.",
        ["frontend/src/routes/SearchPage.tsx:167-187", "frontend/src/routes/ExplorePage.tsx:125-146"],
        [
            "Add AbortController support or adopt a query library with cancellation, caching, deduplication, and stale-response protection.",
            "Fall back to flat search only for a recognized unsupported-endpoint response.",
            "Normalize HTTP errors into authentication, quota, validation, timeout, and outage states.",
        ],
    )
    add_issue(
        doc,
        "Video-player state is unsafe across route changes",
        "P1",
        "YouTubePlayer captures start once, does not reset ready when videoId changes, and reuses mutable player state while destroying and recreating the iframe. Direct navigation between two VOD routes can initialize the second VOD with the first timestamp or seek a destroyed player. Script failures have no visible error state.",
        ["frontend/src/components/YouTubePlayer.tsx:52-109"],
        [
            "Reset readiness and pending state on videoId and use the current start value for each player instance.",
            "Centralize the YouTube API loader in a shared singleton promise and expose loading/error states.",
            "Add a rerender test that changes both videoId and start.",
        ],
    )
    add_issue(
        doc,
        "Authentication initialization produces an unhandled rejection",
        "P1",
        "The /auth/me chain has finally but no catch. The matching error test is skipped because it triggers unhandled rejection warnings. A network failure therefore looks like an anonymous session while also producing an unhandled browser error. Logout suppresses server failure and clears local state even when the session cookie may remain valid.",
        ["frontend/src/services/auth.tsx:25-48", "frontend/src/tests/auth.test.tsx:86-104"],
        [
            "Represent auth as loading, authenticated, anonymous, or error.",
            "Restore the skipped network-error test.",
            "Only present logout as successful after confirmation, or show a retryable failure.",
        ],
    )
    add_issue(
        doc,
        "Route-level not-found and error handling are missing",
        "P1",
        "The router has no wildcard route and no errorElement. Unknown URLs render an empty layout, and unexpected render errors can blank the application. The browser suite expects a 404 page that does not exist. The /login page also still says authentication is coming soon despite working Google and Twitch OAuth entry points.",
        ["frontend/src/main.tsx:25-69", "frontend/src/routes/LoginPage.tsx:1-8"],
        [
            "Add a wildcard 404 page and route-level error boundary with retry and home actions.",
            "Provide a distinct missing-VOD state.",
            "Replace or remove the obsolete login placeholder.",
        ],
    )
    add_issue(
        doc,
        "Admin authorization is not reflected in the frontend contract",
        "P1",
        "The current-user response does not expose a role or capability set, and AdminLayout allows any authenticated user to render the admin shell before backend requests return 403. Backend enforcement prevents privilege escalation, but normal users can reach a broken-looking administrative experience.",
        ["app/routes/auth.py:87-102", "frontend/src/routes/admin/AdminLayout.tsx:5-11"],
        [
            "Return a safe role or capability set from /auth/me.",
            "Require the admin capability before rendering the shell and provide a proper 403 page.",
            "Prefer explicit capabilities to frontend email or domain heuristics.",
        ],
    )
    add_issue(
        doc,
        "Dependency advisories require immediate triage",
        "P1",
        "The current lockfile audit reports high-severity production findings, including react-router-dom 7.9.4, plus critical and high findings in development/build tooling. Several Router advisories concern SSR/RSC paths this SPA does not use, so applicability must be assessed, but the versions should still be updated promptly.",
        ["frontend/package.json", "frontend/package-lock.json"],
        [
            "Upgrade production and toolchain dependencies in a bounded change.",
            "Record which advisories are reachable in this SPA and which are build-time only.",
            "Run router, deep-link, authentication, production-build, and browser regressions after updating.",
        ],
    )

    add_heading(doc, "Performance and Scalability", 1)
    add_heading(doc, "Route bundle", 2)
    add_body(doc, "All public and admin routes are eagerly imported. The production build emits a single 460.42 KB JavaScript asset (130.90 KB gzip) plus 94.14 KB CSS (13.11 KB gzip). Public visitors therefore download large administrative screens they may never use.")
    add_bullet(doc, "Use route-level lazy loading for public pages and especially for the 853-line metadata editor, 588-line period editor, dashboard/chart code, and other admin pages.")
    add_bullet(doc, "Add an enforceable compressed-size budget per entry chunk rather than only warning on a 500 KB raw bundle.")

    add_heading(doc, "Transcript rendering", 2)
    add_body(doc, "VideoPage polls playback every 750 ms and stores currentMs at the route root. That can rerender hundreds of transcript blocks continuously. Match calculation also performs hits.find() for every segment, and smooth scrolling may restart on each polling update.")
    add_bullet(doc, "Isolate playback time in a small subscribed component or store.")
    add_bullet(doc, "Memoize transcript blocks and update only the previously and currently active blocks.")
    add_bullet(doc, "Pre-index hits by timestamp or segment instead of repeated nested scans.")
    add_bullet(doc, "Use content-visibility: auto or virtualization/windowing for long transcripts.")
    add_bullet(doc, "Throttle automatic scrolling and avoid overlapping smooth-scroll animations.")

    add_heading(doc, "Data fetching", 2)
    add_body(doc, "SearchPage requests the complete Explore intelligence response merely to obtain suggested searches. Archive summary, period lists, topic data, and metadata dictionaries are repeatedly fetched without a shared freshness model.")
    add_bullet(doc, "Add a lightweight suggestions endpoint or share cached queries.")
    add_bullet(doc, "Define freshness, retry, and cancellation policies for archive summary, Explore periods, topics, and dictionaries.")

    add_heading(doc, "Maintainability hotspots", 2)
    add_body(doc, "Several route components combine networking, state machines, data conversion, mutations, and presentation:")
    for ref in [
        "AdminVideoMetadata.tsx — 853 lines",
        "VideoPage.tsx — 724 lines",
        "ExplorePage.tsx — 660 lines",
        "AdminArchivePeriods.tsx — 588 lines",
        "SearchPage.tsx — 485 lines",
    ]:
        add_bullet(doc, ref)
    add_body(doc, "Split these by feature responsibility: query hooks, URL-state adapters, view models, mutations, and presentational sections. VideoPage especially needs a playback controller hook and a transcript-navigation model. Add a root verify command that runs tests, coverage, lint, formatting, type-check, build, and selected browser smoke tests.")

    add_heading(doc, "Accessibility and UI Consistency", 1)
    findings = [
        ("Mobile menu focus", "The collapsed menu uses aria-hidden but leaves links in the DOM and keyboard tab order. Conditionally render it or use hidden/inert; support Escape and return focus to the menu button.", "frontend/src/routes/AppLayout.tsx:173-290"),
        ("Nested landmarks", "Search and Explore place additional main elements inside AppLayout's main. Replace them with section or div elements.", "frontend/src/routes/SearchPage.tsx:296; frontend/src/routes/ExplorePage.tsx:397"),
        ("Incorrect tab semantics", "Explore declares role=tablist but uses pressed buttons without role=tab or arrow-key behavior. Implement tabs completely or use a labelled button group.", "frontend/src/routes/ExplorePage.tsx:327-346"),
        ("Unannounced selection", "Transcript layout buttons do not expose the active mode with aria-pressed.", "frontend/src/routes/VideoPage.tsx:501-511"),
        ("Touch targets", "Several shared controls are 38-40 px despite documentation claiming a 44 px minimum.", "frontend/src/index.css:148-181"),
        ("Transitions", "transition-all is used repeatedly; enumerate only the intended properties.", "frontend/src/index.css:153-181"),
        ("Image layout stability", "Several thumbnail images omit explicit width and height.", "frontend/src/components/archive/VideoCard.tsx:25-30"),
        ("Form metadata", "Form fields generally lack meaningful name and autocomplete attributes.", "frontend/src/components/archive/SearchFiltersPanel.tsx:45-84"),
        ("Automated accessibility", "Axe dependencies are installed but not used. Add assertions for layout, search, Explore, Saved, Topic, and Video states.", "frontend/package.json"),
    ]
    for title_text, body, ref in findings:
        add_bullet(doc, f"{title_text}: {body}", bold_prefix=f"{title_text}:")
        add_code_ref(doc, ref)

    add_heading(doc, "Missing or Incomplete Product Features", 1)
    add_body(doc, "Against the product direction recorded in docs/frontend-overhaul.md, the largest gaps are:")
    for item in [
        "Topic history has summary statistics but no topic-over-time visualization.",
        "There is no citation-backed opinion-over-time experience.",
        "Episode pages lack related episodes and most-quoted moments.",
        "Every-mention collections cannot be exported as a list or playlist.",
        "Timeline exists but is absent from primary navigation.",
        "Explore period selections are not reflected in the URL, so they cannot be bookmarked or shared.",
        "PWA and offline behavior are intentionally disabled while PWA assets and claims remain.",
        "Copy and save actions often lack success or failure feedback.",
    ]:
        add_bullet(doc, item)
    add_callout(
        doc,
        "Regression coverage gap",
        "The previously observed blank topic-detail and missing people/tag behavior appears to be the target of commit 873a78b, but frontend tests only use mocked responses. Add one seeded browser flow from an Explore topic card to a nonempty Topic page containing statistics, first/latest mentions, related topics, and grouped results.",
        color=GOLD,
        fill="FFF8E1",
    )

    add_heading(doc, "Documentation Review", 1)
    add_heading(doc, "README and browser-test documentation", 2)
    add_body(doc, "The README claims a pricing/upgrade flow and upgrade interstitial, but no pricing route or frontend billing implementation exists. Billing E2E tests navigate to /pricing, while job tests expect job-creation screens and /jobs/:id routes that are also absent.")
    add_body(doc, "Browser specifications also use /videos/:id instead of /v/:id, omit the required { user: ... } auth envelope, expect automatic authenticated redirects, expect a nonexistent 404 page, and return search fields such as results and segment_text instead of hits and snippet.")
    add_body(doc, "The published '255 E2E tests' figure describes scenario/browser combinations rather than 255 current valid workflows. The PR-critical CI runs the particularly stale auth, job, and search specifications.")
    add_bullet(doc, "Rewrite the browser suite around current archive workflows: anonymous search, filters, Explore, topic detail, VOD playback and deep links, local and synchronized saves, exports, admin denial, and mobile navigation.")

    add_heading(doc, "Architecture documentation", 2)
    add_body(doc, "docs/development/architecture.md describes React 18 and Axios. The current frontend uses React 19 and Ky. It also documents frontend job submission and polling that do not exist.")

    add_heading(doc, "Design-system documentation", 2)
    add_body(doc, "docs/DESIGN_SYSTEM.md describes the former blue/orange Satoshi SaaS theme. The implementation now uses Alegreya/Atkinson, dark editorial surfaces, a lime accent, and a purple CTA. The token table and referenced pricing surfaces are obsolete.")

    add_heading(doc, "Accessibility and PWA documentation", 2)
    add_body(doc, "docs/ACCESSIBILITY.md claims WCAG 2.1 AA compliance, Escape-to-close, a slash search shortcut, 44 px targets, working offline support/background sync, and roughly 101 KB gzip JavaScript. Several claims are false. Most significantly, main.tsx unregisters every service worker and deletes all caches while documentation still describes a working offline PWA.")
    add_bullet(doc, "Change compliance language to 'targets WCAG 2.2 AA' until independently audited.")
    add_bullet(doc, "Record the browser and assistive-technology combinations actually tested.")
    add_bullet(doc, "State PWA status explicitly and either restore it or remove the obsolete assets and installation claims.")

    add_heading(doc, "Developer testing documentation", 2)
    add_body(doc, "docs/development/testing.md documents npm run test:watch, which does not exist. Frontend CI documentation should also state that lint and formatting are blocking checks; the current local failures are not merely informational warnings.")

    add_heading(doc, "Documentation Improvement Model", 2)
    add_bullet(doc, "Separate shipped behavior, planned work, and historical implementation notes into clearly labelled documents.")
    add_bullet(doc, "Add a frontend route and capability inventory maintained beside the router.")
    add_bullet(doc, "Generate or validate frontend API types against the OpenAPI schema to prevent envelope and field drift.")
    add_bullet(doc, "Create one canonical developer workflow with setup, verify, browser smoke, and troubleshooting commands.")
    add_bullet(doc, "Treat accessibility and performance figures as measured evidence with a date and environment, not permanent claims.")

    add_heading(doc, "Verification Results", 1)
    add_verification_table(doc)
    add_body(doc, "The worktree remained unchanged apart from the pre-existing untracked .codex directory. The Docker-backed browser suite was not started because this was a read-only review and the Playwright configuration creates external service state.")

    add_heading(doc, "Recommended Implementation Order", 1, page_break=True)
    steps = [
        "Eliminate unsafe HTML rendering and update vulnerable dependencies.",
        "Restore a green frontend gate: lint, formatting, the auth error test, and current browser smoke tests.",
        "Fix user-facing contract breaks: anonymous Saved access, preserved search filters, 404/error boundaries, and auth/admin capabilities.",
        "Add request cancellation and correct YouTube-player route-change behavior.",
        "Split routes and isolate long-transcript playback updates.",
        "Reconcile README, architecture, design-system, accessibility, PWA, and E2E documentation.",
        "Implement product-direction gaps such as topic timelines, related episodes, cited opinion history, and exportable mention collections.",
    ]
    for step in steps:
        add_numbered(doc, step)

    add_heading(doc, "Acceptance Criteria for the Next Review", 1)
    criteria = [
        "No user-controlled or externally sourced string reaches dangerouslySetInnerHTML unsanitized.",
        "Anonymous users can save and reopen local moments.",
        "Every serialized search filter is honored or rejected explicitly.",
        "Changing query, period, or video cannot display stale prior responses.",
        "Unknown routes and missing videos have tested recovery pages.",
        "Admin routes render only for users with the corresponding capability.",
        "Unit tests, coverage, type-check, lint, formatting, build, and a current browser smoke set all pass from one command.",
        "Public-route JavaScript excludes admin bundles and meets an explicit compressed-size budget.",
        "Accessibility and PWA documentation matches tested behavior.",
        "README and architecture documentation match the current router, packages, and API envelopes.",
    ]
    for item in criteria:
        add_bullet(doc, item)

    add_heading(doc, "Reference", 1)
    add_body(doc, "Vercel Web Interface Guidelines: https://raw.githubusercontent.com/vercel-labs/web-interface-guidelines/main/command.md")

    # Document metadata.
    props = doc.core_properties
    props.title = "HasAnAra Web Frontend Quality Review"
    props.subject = "Code quality, functionality, performance, accessibility, scalability, and documentation assessment"
    props.author = "Codex"
    props.keywords = "HasAnAra, frontend, React, quality review, accessibility, performance, documentation"
    props.comments = "Generated from a read-only repository review."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
